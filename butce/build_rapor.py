# -*- coding: utf-8 -*-
"""
Bütçe sunum özeti (PDF). Hesaplanmış Excel'den (LibreOffice ile yeniden hesaplanmış) verileri okur,
HTML üretir ve Chromium (node playwright) ile PDF'e basar.
  python3 build_rapor.py URF_1Etap_MEP_Butce_2026-10-06.xlsx veri/rfq_firmalar.csv
"""
import sys, os, csv, html, json, subprocess, collections
import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else "URF_1Etap_MEP_Butce_2026-10-06.xlsx"
FIRMS = sys.argv[2] if len(sys.argv) > 2 else "veri/rfq_firmalar.csv"
OUT_HTML = "URF_1Etap_MEP_Butce_Sunum_2026-10-06.html"
OUT_PDF = "URF_1Etap_MEP_Butce_Sunum_2026-10-06.pdf"

wb = openpyxl.load_workbook(XLSX, data_only=True)
O = wb["Ozet"]
e = lambda v: "" if v is None else f"{v:,.0f}".replace(",", ".")
pc = lambda v: "" if v in (None, "") else f"%{v*100:.0f}"
esc = lambda t: html.escape(str(t if t is not None else ""))

# ---- disiplin tablosu
dis = []
for r in O.iter_rows(min_row=5, max_row=10, values_only=True):
    dis.append(dict(ad=r[0], kapsam=r[1], kalem=r[2], risk=r[3], oner=r[4], usd=r[5], tl=r[6], pay=r[7], rfq=r[8],
                    tek=r[10], ems=r[11], tah=r[12]))
TOP = dis[-1]; DIS = dis[:-1]

# ---- alt kırılım
alt = []
start = None
for i, r in enumerate(O.iter_rows(min_row=1, max_row=60, values_only=True), 1):
    if r[0] == "ALT KIRILIM (disiplin içi)":
        start = i + 2
    if start and i >= start:
        if not r[0] or not str(r[0])[0].isdigit():
            break
        alt.append(dict(d=r[0], ad=r[1], sira=r[2], kalem=r[3], oner=r[5], rfq=r[6], tek=r[10], ems=r[11], tah=r[12]))

# ---- vurgu blokları (Özet A sütunu metinlerinden)
rows = list(O.iter_rows(min_row=1, max_row=120, max_col=7, values_only=True))
def block(title_prefix, ncols=7):
    out, on = [], False
    for r in rows:
        if r[0] and str(r[0]).startswith(title_prefix):
            on = True; continue
        if on:
            if r[0] and str(r[0]).startswith("▶") or (r[0] and str(r[0]).startswith("CETVEL")):
                break
            out.append(r)
    return out
kd = block("▶ KAPSAM DIŞI KALEMLER")
de = block("▶ KAPSAM DIŞI EKİPMANLAR")
rb = block("▶ İHALE GENELİ")

# ---- teklif değerlendirme
T = wb["Teklif_Degerlendirme"]
tek = []
for r in T.iter_rows(min_row=5, max_row=13, values_only=True):
    if r[1]:
        tek.append(r)

# ---- disiplin bazında en büyük kalemler
SH = {"1 ELEKTRİK": ["Elektrik_AG", "Haberlesme"], "2 MEKANİK": ["Mekanik"], "3 ASANSÖR": ["Asansor"],
      "4 SİNYALİZASYON": ["Sinyal"], "5 ELEKTRİFİKASYON": ["Cer_Guc", "Katener"]}
top_items = {}
for d, shs in SH.items():
    L = []
    for sh in shs:
        ws = wb[sh]
        sira = None
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r[0]:
                sira = r[0]; continue
            if r[2] and r[3] and isinstance(r[9], (int, float)) and r[10]:
                L.append((r[9], sira, r[2], r[3], r[7], r[8], r[10], (r[11] or "")))
    L.sort(key=lambda x: -x[0])
    top_items[d] = L[:8]

# ---- firmalar
firms = []
if os.path.exists(FIRMS):
    firms = list(csv.DictReader(open(FIRMS, encoding="utf-8-sig")))

# ---- Varsayımlar
V = wb["Varsayimlar"]
usd = V["B4"].value; tl = V["B5"].value

TYPE_COL = {"2026 TEKLİF": "#70AD47", "GEÇMİŞ TEKLİF UYARLAMA": "#5B9BD5", "MÜHENDİSLİK TAHMİNİ": "#E2A800", "KAPSAM DIŞI (0)": "#A6A6A6"}

# ---- SVG yığılmış çubuk grafik
def bar_svg():
    W, H, lw, bh, gap = 860, 265, 190, 30, 14
    mx = max((d["tek"] or 0) * d["kalem"] + (d["ems"] or 0) * d["kalem"] + (d["tah"] or 0) * d["kalem"] for d in DIS)
    sc = (W - lw - 120) / mx
    s = [f'<svg viewBox="0 0 {W} {H}" width="100%" xmlns="http://www.w3.org/2000/svg" font-family="Arial" font-size="12">']
    y = 10
    for d in DIS:
        x = lw
        s.append(f'<text x="{lw-8}" y="{y+bh/2+4}" text-anchor="end" fill="#1F2937" font-weight="bold">{esc(d["ad"])}</text>')
        for key, col in (("tek", "#70AD47"), ("ems", "#5B9BD5"), ("tah", "#E2A800")):
            v = (d[key] or 0) * d["kalem"]
            w = v * sc
            if w > 0:
                s.append(f'<rect x="{x:.1f}" y="{y}" width="{w:.1f}" height="{bh}" fill="{col}" rx="2"/>')
                if w > 55:
                    s.append(f'<text x="{x+w/2:.1f}" y="{y+bh/2+4}" text-anchor="middle" fill="#fff" font-size="11">{v/1e6:.2f}</text>')
            x += w
        s.append(f'<text x="{x+6:.1f}" y="{y+bh/2+4}" fill="#1F2937" font-weight="bold">{d["oner"]/1e6:.2f} M€</text>')
        y += bh + gap
    ly = y + 4
    for i, (lab, col) in enumerate((("2026 teklifi", "#70AD47"), ("Geçmiş teklif uyarlaması", "#5B9BD5"), ("Mühendislik tahmini", "#E2A800"))):
        lx = lw + i * 200
        s.append(f'<rect x="{lx}" y="{ly}" width="12" height="12" fill="{col}"/><text x="{lx+18}" y="{ly+10}" fill="#374151">{lab}</text>')
    s.append(f'<text x="{lw}" y="{ly+30}" fill="#6B7280" font-size="10">Çubuk içi: kalem bazlı M€ · çubuk sağı: risk dahil önerilen bütçe</text>')
    s.append("</svg>")
    return "".join(s)

# ---- HTML
css = """
@page { size: A4; margin: 14mm 13mm 16mm 13mm; }
* { box-sizing: border-box; }
body { font-family: Arial, Helvetica, sans-serif; color: #1F2937; font-size: 10.5px; line-height: 1.38; margin: 0; }
h1 { font-size: 26px; color: #1F4E78; margin: 0 0 6px; }
h2 { font-size: 16px; color: #fff; background: #1F4E78; padding: 7px 10px; margin: 18px 0 10px; border-radius: 3px; }
h3 { font-size: 12.5px; color: #1F4E78; margin: 14px 0 6px; border-bottom: 1.5px solid #BDD7EE; padding-bottom: 3px; }
.page { page-break-after: always; }
.cover { height: 260mm; display: flex; flex-direction: column; justify-content: center; }
.cover .band { background: #1F4E78; color: #fff; padding: 28px 30px; border-radius: 6px; }
.cover .band h1 { color: #fff; font-size: 30px; }
.cover .sub { font-size: 14px; opacity: .9; }
.meta { margin-top: 26px; font-size: 12px; color: #374151; }
.meta td { padding: 4px 14px 4px 0; }
.kpis { display: grid; grid-template-columns: repeat(4, 1fr); gap: 8px; margin: 8px 0 12px; }
.kpi { border: 1px solid #BDD7EE; border-radius: 6px; padding: 9px 10px; background: #F5F9FD; }
.kpi .l { font-size: 9.5px; color: #4B5563; text-transform: uppercase; letter-spacing: .3px; }
.kpi .v { font-size: 18px; font-weight: bold; color: #1F4E78; margin-top: 3px; }
.kpi .s { font-size: 9.5px; color: #6B7280; }
table { width: 100%; border-collapse: collapse; margin: 4px 0 8px; }
th { background: #1F4E78; color: #fff; font-size: 9.5px; padding: 5px 6px; text-align: left; vertical-align: bottom; }
td { border-bottom: 1px solid #E5E7EB; padding: 4px 6px; vertical-align: top; font-size: 9.8px; }
td.n, th.n { text-align: right; white-space: nowrap; }
tr.tot td { font-weight: bold; background: #E2EFDA; border-top: 1.5px solid #70AD47; }
tr.sub td { background: #F9FAFB; color: #374151; }
.hl { border: 2px solid #C55A11; border-radius: 6px; padding: 8px 10px; margin: 10px 0; background: #FFF7F0; }
.hl .t { color: #fff; background: #C55A11; margin: -8px -10px 8px; padding: 6px 10px; font-weight: bold; font-size: 11.5px; border-radius: 4px 4px 0 0; }
.note { color: #4B5563; font-size: 9.5px; }
.pill { display: inline-block; padding: 1px 6px; border-radius: 8px; color: #fff; font-size: 8.5px; white-space: nowrap; }
.two { display: grid; grid-template-columns: 1fr 1fr; gap: 12px; }
.small td, .small th { font-size: 8.8px; padding: 3px 5px; }
.step { display: grid; grid-template-columns: 26px 1fr; gap: 6px; margin: 4px 0; }
.step .no { background: #1F4E78; color: #fff; border-radius: 50%; width: 22px; height: 22px; text-align: center; line-height: 22px; font-weight: bold; }
.footer { position: fixed; bottom: -10mm; left: 0; right: 0; font-size: 8.5px; color: #9CA3AF; text-align: right; }
"""
P = []
P.append(f"""<div class="page cover"><div class="band">
<div class="sub">ŞANLIURFA HALİLİYE–EYYÜBİYE HAFİF RAYLI SİSTEM HATTI 1. AŞAMA YAPIM İŞİ · İKN 2026/1327584</div>
<h1>MEP Bütçe Çalışması – Sunum Özeti</h1>
<div class="sub">Elektrik · Mekanik · Asansör · Sinyalizasyon · Elektrifikasyon</div></div>
<table class="meta"><tr><td><b>Teklif sahibi</b></td><td>ÖZVER İNŞAAT A.Ş.</td></tr>
<tr><td><b>Hazırlayan</b></td><td>MEP Center – Serhat Çelik</td></tr>
<tr><td><b>Tarih</b></td><td>06.10.2026 (ihale tarihi 13.10.2026)</td></tr>
<tr><td><b>Fiyat esası</b></td><td>EUR, KDV hariç, Şanlıurfa şantiye teslim, montaj + test + devreye alma dahil (taşeron fiyatı)</td></tr>
<tr><td><b>Kur</b></td><td>1 EUR = {usd:.4f} USD = {tl:.3f} TL (ECB, 02.10.2026)</td></tr>
<tr><td><b>Ekler</b></td><td>URF_1Etap_MEP_Butce_2026-10-06.xlsx – kalem kalem formüllü çalışma (1.625 kalem)</td></tr></table>
</div>""")

emsal_pay = (TOP["tek"] or 0) + (TOP["ems"] or 0)
P.append(f"""<div class="page"><h2>1. Yönetici Özeti</h2>
<div class="kpis">
<div class="kpi"><div class="l">Toplam MEP önerilen bütçe</div><div class="v">{TOP['oner']/1e6:.2f} M€</div><div class="s">risk payı dahil · kalem bazlı {TOP['kalem']/1e6:.2f} M€</div></div>
<div class="kpi"><div class="l">USD / TL karşılığı</div><div class="v">{TOP['usd']/1e6:.1f} M$</div><div class="s">≈ {TOP['tl']/1e9:.2f} milyar TL</div></div>
<div class="kpi"><div class="l">Teklif + emsale dayanan pay</div><div class="v">%{emsal_pay*100:.0f}</div><div class="s">2026 teklif %{(TOP['tek'] or 0)*100:.0f} · geçmiş teklif %{(TOP['ems'] or 0)*100:.0f}</div></div>
<div class="kpi"><div class="l">Kalem kalem analiz</div><div class="v">1.625 kalem</div><div class="s">1.183 geçmiş teklif satırı · 28 proje · 103 firma</div></div>
</div>
<table><tr><th>Disiplin</th><th>Kapsam</th><th class="n">Kalem bazlı €</th><th class="n">Risk €</th><th class="n">Önerilen €</th><th class="n">Pay</th><th class="n">Teklif</th><th class="n">Emsal</th><th class="n">Tahmin</th></tr>
""" + "".join(f"<tr><td><b>{esc(d['ad'])}</b></td><td>{esc(d['kapsam'])}</td><td class='n'>{e(d['kalem'])}</td><td class='n'>{e(d['risk'])}</td><td class='n'><b>{e(d['oner'])}</b></td><td class='n'>{pc(d['pay'])}</td><td class='n'>{pc(d['tek'])}</td><td class='n'>{pc(d['ems'])}</td><td class='n'>{pc(d['tah'])}</td></tr>" for d in DIS)
+ f"<tr class='tot'><td>TOPLAM MEP</td><td></td><td class='n'>{e(TOP['kalem'])}</td><td class='n'>{e(TOP['risk'])}</td><td class='n'>{e(TOP['oner'])}</td><td class='n'>%100</td><td class='n'>{pc(TOP['tek'])}</td><td class='n'>{pc(TOP['ems'])}</td><td class='n'>{pc(TOP['tah'])}</td></tr></table>"
+ f"<h3>Disiplin bazında bütçe ve fiyat kaynağı</h3>{bar_svg()}"
+ """<h3>Öne çıkanlar</h3><ul>
<li><b>Kalem kalem:</b> 5 disiplinin teklif talebi cetvellerindeki tüm alt kalemler miktar × birim fiyat olarak fiyatlandı; her satırda fiyatın kaynağı (2026 teklif / geçmiş teklif emsali ID'si / tahmin gerekçesi) yazılıdır.</li>
<li><b>Net metraj:</b> Tedarikçilere giden cetvellerdeki bilinçli metraj payları (%10 fire, dağıtık malzemede %30) bütçeden geri alındı; idarenin 05.10.2026 cevapları uygulandı.</li>
<li><b>Gelen teklifler bütçeyi doğruluyor:</b> Asansör (3 teklif) ±%16, AVLS (Mukan) +%3, saat ve yolcu bilgilendirme (ON Elektronik) ±%6 bandında.</li>
<li><b>Elektrifikasyon en büyük kalem:</b> 5 cer merkezi, OG ring ve TEİAŞ bağlantı kabloları (transe dahil) ve katener ile toplamın ~%43'ü.</li>
</ul></div>""")

# Disiplin detay sayfası
P.append('<div class="page"><h2>2. Disiplin Bazında Kırılım</h2>')
P.append("<table><tr><th>Disiplin</th><th>Alt kalem</th><th>Cetvel sıraları</th><th class='n'>Kalem bazlı €</th><th class='n'>Önerilen €</th><th class='n'>RFQ metrajıyla €</th></tr>"
         + "".join(f"<tr><td><b>{esc(a['d'])}</b></td><td>{esc(a['ad'])}</td><td class='note'>{esc(a['sira'])}</td><td class='n'>{e(a['kalem'])}</td><td class='n'><b>{e(a['oner'])}</b></td><td class='n'>{e(a['rfq'])}</td></tr>" for a in alt)
         + "</table><p class='note'>RFQ metrajıyla: tedarikçilere gönderilen (şişirmeli, idare düzeltmesiz) miktarlarla aynı birim fiyatlardan hesaplanan tutar – net bütçenin ne kadar temkinli olduğunu gösterir.</p>")
for d, L in top_items.items():
    P.append(f"<h3>{esc(d)} – en büyük kalemler</h3><table class='small'><tr><th>Sıra</th><th>Kalem</th><th>Birim</th><th class='n'>Net miktar</th><th class='n'>Birim fiyat €</th><th class='n'>Tutar €</th><th>Fiyat türü</th></tr>")
    for t, sira, kal, bir, q, bf, tur, src in L:
        P.append(f"<tr><td>{esc(sira)}</td><td>{esc(kal[:110])}</td><td>{esc(bir)}</td><td class='n'>{(q or 0):,.1f}</td><td class='n'>{e(bf)}</td><td class='n'><b>{e(t)}</b></td><td><span class='pill' style='background:{TYPE_COL.get(tur,'#999')}'>{esc(tur)}</span></td></tr>")
    P.append("</table>")
P.append("</div>")

# Teklifler
P.append('<div class="page"><h2>3. Gelen Tekliflerin Değerlendirmesi</h2>')
P.append("<table><tr><th>Disiplin</th><th>Firma</th><th>Tarih</th><th>Kapsam</th><th class='n'>Teklif €</th><th class='n'>Bütçe karşılığı €</th><th class='n'>Fark</th><th>Değerlendirme</th></tr>")
for r in tek:
    if r[0] == "Bilgi – MEP dışı":
        continue
    fark = "" if not isinstance(r[8], (int, float)) else f"{r[8]*100:+.0f}%"
    P.append(f"<tr><td>{esc(r[0])}</td><td><b>{esc(r[1])}</b></td><td>{esc(r[2])}</td><td>{esc(r[3])}</td><td class='n'>{e(r[6])}</td><td class='n'>{e(r[7])}</td><td class='n'><b>{fark}</b></td><td class='note'>{esc(r[13])}</td></tr>")
P.append("</table><h3>Teknik ve ticari notlar</h3><table class='small'><tr><th>Firma</th><th>Teknik uygunluk / sapmalar</th><th>Ticari şartlar</th><th>Bütçede kullanımı</th></tr>")
for r in tek:
    if r[0] == "Bilgi – MEP dışı":
        continue
    P.append(f"<tr><td><b>{esc(r[1])}</b></td><td>{esc(r[10])}</td><td>{esc(r[11])}</td><td>{esc(r[12])}</td></tr>")
P.append("</table>")
if "PointLink_Kiyas" in wb.sheetnames:
    PLs = wb["PointLink_Kiyas"]
    P.append("<h3>Point Link / CASCO teklifi – kalem kalem kıyas (sinyalizasyon + araç takip)</h3><table class='small'><tr><th>Sıra</th><th>Kalem</th><th class='n'>PL teklif €</th><th class='n'>PL idareye göre uyarlanmış €</th><th class='n'>Bütçemiz €</th><th class='n'>Fark</th></tr>")
    for r in PLs.iter_rows(min_row=5, max_row=16, values_only=True):
        if not r[1]:
            continue
        lab = str(r[1])
        if lab.startswith("BİRİM"):
            break
        cls = " class='tot'" if lab.startswith(("Ara toplam", "GENEL")) else ""
        fark = f"{r[6]*100:+.0f}%" if isinstance(r[6], (int, float)) else ""
        P.append(f"<tr{cls}><td>{esc(r[0] or '')}</td><td>{esc(lab)}</td><td class='n'>{e(r[2])}</td><td class='n'>{e(r[3])}</td><td class='n'>{e(r[5]) if isinstance(r[5],(int,float)) else ''}</td><td class='n'><b>{fark}</b></td></tr>")
    P.append("</table><p class='note'>Point Link fiyatları kontrol kabinleri (SKP 166.750 €, TSP 78.200 €), yazılım (299.000 €), tasarım/ISA (540.000 €) ve test-devreye alma (360.000 €) kalemlerinde yüksektir; makas motoru (7.950 €) ve aks sayacı (5.340 €) bütçemizin altındadır. Teklif Fwd ile ulaşmış, yapay zekâ imzalıdır ve orijinal mail mepcenter Gmail'de yoktur – kaynak teyidi gerekir. Bütçede kullanılmamış, üst sınır senaryosu olarak raporlanmıştır.</p>")

# Kapsam dışı + Rayba (vurgu)
P.append('<h2>4. Kapsam Dışı Kalemler ve İhale Geneli Bilgi</h2>')
P.append('<div class="hl"><div class="t">▶ İdare cevaplarıyla (05.10.2026) bütçeden düşülen kalemler – RFQ birim fiyatlarıyla değeri</div><table class="small"><tr><th>Disiplin sayfası</th><th>Düşülen başlıca kalemler</th><th class="n">Değer €</th></tr>')
for r in kd:
    if r[0] in (None, "Disiplin sayfası"):
        continue
    if str(r[0]).startswith("Ayrıca"):
        break
    cls = " class='tot'" if r[0] == "Toplam düşülen" else ""
    P.append(f"<tr{cls}><td>{esc(r[0])}</td><td>{esc(r[1] or '')}</td><td class='n'>{e(r[4])}</td></tr>")
P.append("</table></div>")
P.append('<div class="hl"><div class="t">▶ Kapsam dışı ekipmanlar – yüklenici kapsamındaki depo ekipmanları (MEP paketlerinde yok) – bilgi amaçlı tahmin</div><table class="small"><tr><th>Ekipman</th><th>Açıklama</th><th class="n">Miktar</th><th class="n">Birim €</th><th class="n">Tutar €</th></tr>')
for r in de:
    if r[0] in (None, "Ekipman") or str(r[0]).startswith("İdare cevabı"):
        continue
    cls = " class='tot'" if str(r[0]).startswith("Toplam") else ""
    P.append(f"<tr{cls}><td>{esc(r[0])}</td><td>{esc(r[1] or '')}</td><td class='n'>{esc(r[2] or '')}</td><td class='n'>{e(r[3]) if isinstance(r[3],(int,float)) else ''}</td><td class='n'>{e(r[4])}</td></tr>")
P.append("</table><p class='note'>Veritabanında emsal yoktur; teklif alınana kadar mertebe göstergesidir.</p></div>")
P.append('<div class="hl"><div class="t">▶ İhale geneli – MEP dışı gelen teklif: RAYBA YAPI (demiryolu üstyapı) – MEP toplamına dahil değil</div><table class="small"><tr><th>Teklif</th><th class="n">EUR</th><th class="n">USD</th><th class="n">TL</th></tr>')
for r in rb:
    if r[0] == "Rayba Yapı":
        P.append(f"<tr><td>{esc(r[1])}</td><td class='n'><b>{e(r[4])}</b></td><td class='n'>{e(r[5])}</td><td class='n'>{e(r[6])}</td></tr>")
P.append("</table><p class='note'>42×R50 + 2×R100 makas + 3 kruvazman (= 56 makas) sinyal/katener kabulünü teyit ediyor. S-LINE kapsülleri hat döşeme fiyatında olabilir (çift sayım kontrolü); beton/donatı hariç; 31.12.2026'ya kadar geçerli.</p></div></div>")

# Yöntem
P.append("""<div class="page"><h2>5. Birim Fiyatlar Nasıl Belirlendi?</h2>
<p>Bütçedeki her birim fiyat, izlenebilir bir kaynağa dayanır. Öncelik sırası: <span class="pill" style="background:#70AD47">2026 TEKLİF</span> →
<span class="pill" style="background:#5B9BD5">GEÇMİŞ TEKLİF UYARLAMA</span> → <span class="pill" style="background:#E2A800">MÜHENDİSLİK TAHMİNİ</span>.</p>
<h3>Geçmiş teklif veritabanı</h3>
<div class="two"><div><table class="small"><tr><th>Gösterge</th><th class="n">Veritabanı</th><th class="n">Bu bütçede kullanılan</th></tr>
<tr><td>Fiyat satırı</td><td class="n">6.809</td><td class="n">1.183</td></tr><tr><td>Proje</td><td class="n">42</td><td class="n">28</td></tr>
<tr><td>Firma</td><td class="n">222</td><td class="n">103</td></tr><tr><td>Tramvay / LRT satırı</td><td class="n">1.352</td><td class="n">535</td></tr>
<tr><td>Metro satırı</td><td class="n">3.273</td><td class="n">509</td></tr><tr><td>Teklif yılları</td><td class="n">2008–2025</td><td class="n">ağırlıkla 2014–2025</td></tr></table></div>
<div class="note">Kaynak: MEP Center proje arşivindeki raylı sistem projelerinin “MALIYET / GELEN_TEKLIF” klasörlerindeki gerçek tedarikçi teklifleri. Başlıca tramvay referansları: Bursa T2 (Haluk, MET, AKE), İzmir tramvayı (Siemens, Usluel, Hanning & Kahl, CONTEC, Aktif), Kocaeli tramvayı, Konya Alaaddin–Adliye, Eminönü–Alibeyköy (Elektroline), Samsun LRT. Metro emsalleri yalnız tramvay emsali yoksa ve ölçek düzeltmesiyle kullanıldı. Her emsal satırı Excel'deki <b>Emsal_Kaynaklari</b> sayfasında proje, firma, tarih, orijinal fiyat ve kaynak dosyasıyla listelenmiştir.</div></div>
<h3>Fiyatı 2026'ya taşıma zinciri</h3>
<div class="step"><div class="no">1</div><div>Orijinal fiyat (teklifteki para birimi) → teklif günündeki <b>ECB referans kuru</b> ile EUR'ya çevrilir.</div></div>
<div class="step"><div class="no">2</div><div>EUR fiyat, <b>Euro Bölgesi HICP</b> endeksiyle Ağustos 2026'ya taşınır: × HICP(Ağu-2026) ÷ HICP(teklif ayı).</div></div>
<div class="step"><div class="no">3</div><div>Kapsam düzeltmesi: emsal yalnız malzeme ise montaj payı eklenir (kablo +%25–35, armatür/dedektör +%30, ekipman +%8–20).</div></div>
<div class="step"><div class="no">4</div><div>Boyut/teknoloji düzeltmesi: trafo kVA^0,6, VRF kW^0,8; eski elektronikte yaş indirimi; bakır yoğun kalemde bakır düzeltmesi.</div></div>
<div class="step"><div class="no">5</div><div>En az iki bağımsız emsalin <b>medyanı</b> alınır; tek emsalde mevcut tahminle %50 harmanlanır; iç maliyet tahminleri düşük ağırlıklıdır.</div></div>
<h3>Örnek hesap – DC fider hücresi (Cer merkezi, sıra 188)</h3>
<table class="small"><tr><th>Emsal</th><th class="n">Orijinal</th><th class="n">Kur</th><th class="n">EUR (teklif günü)</th><th class="n">HICP katsayısı</th><th class="n">EUR Ağu-2026</th></tr>
<tr><td>Haluk Elektrik – Bursa T2 – 04.06.2015</td><td class="n">102.250 TL</td><td class="n">3,0377</td><td class="n">33.660</td><td class="n">× 1,3262</td><td class="n"><b>44.641</b></td></tr>
<tr><td>MET Mühendislik – Bursa T2 – 08.06.2015</td><td class="n">101.863 TL</td><td class="n">3,0865</td><td class="n">33.003</td><td class="n">× 1,3262</td><td class="n"><b>43.769</b></td></tr>
<tr><td>Kocaeli tramvayı – 2015</td><td class="n"></td><td class="n"></td><td class="n"></td><td class="n"></td><td class="n"><b>74.194</b></td></tr>
<tr class="tot"><td>Bütçe birim fiyatı = medyan</td><td></td><td></td><td></td><td></td><td class="n">44.641 €</td></tr></table>
<h3>Miktarlar ve idare cevapları</h3>
<p>Miktarlar idarenin birim fiyat cetveli, ihale projeleri ve şartnamelerden çıkarılan teklif talebi cetvellerinden alınmıştır. Bütçede net metraj kullanılır. İdarenin 05.10.2026 cevapları uygulanmıştır: 14 araç, tüm makaslar motorlu, ana hatta lente yok, kavşak TSKP kapsam dışı, hat 6,44 km, OG ring 7.964 m, fiber 24/8 core SM, istasyon UPS 20 dk, iç ihtiyaç trafoları Cu/Cu.</p>
<h3>Sınırlamalar</h3><p class="note">Emsal bulunamayan kalemler (toplamın ~%20'si: tasarım, test, yazılım, inşaat payları vb.) mühendislik tahminidir ve teklif geldikçe güncellenecektir. Euro Bölgesi enflasyonu yerel işçilik artışını tam yansıtmayabilir. KDV, gümrük ve bakım hariçtir.</p></div>""")

# Firmalar
if firms:
    P.append('<div class="page"><h2>6. Teklif İstenen Firmalar</h2>')
    cnt = collections.Counter(); cnt_t = collections.Counter(); cnt_d = collections.Counter()
    for f in firms:
        for p_ in (f.get("paketler") or "").replace(",", ";").split(";"):
            if p_.strip(): cnt[p_.strip()] += 1
        cnt_t[f.get("firma_turu") or "Belirsiz"] += 1; cnt_d[f.get("durum") or "CEVAP YOK"] += 1
    P.append("<div class='two'><div><table class='small'><tr><th>Paket</th><th class='n'>Firma</th></tr>" + "".join(f"<tr><td>{esc(k)}</td><td class='n'>{v}</td></tr>" for k, v in cnt.most_common()) + f"<tr class='tot'><td>Toplam firma</td><td class='n'>{len(firms)}</td></tr></table></div>")
    P.append("<div><table class='small'><tr><th>Dönüş durumu</th><th class='n'>Firma</th></tr>" + "".join(f"<tr><td>{esc(k)}</td><td class='n'>{v}</td></tr>" for k, v in cnt_d.most_common()) + "</table>"
             + "<table class='small'><tr><th>Firma türü</th><th class='n'>Firma</th></tr>" + "".join(f"<tr><td>{esc(k)}</td><td class='n'>{v}</td></tr>" for k, v in cnt_t.most_common()) + "</table></div></div>")
    order = ["ELEKTRIK", "MEKANIK", "ASANSOR", "SINYALIZASYON", "KATENER"]
    def key(f):
        ps = (f.get("paketler") or "").upper()
        return (min([order.index(p) for p in order if p in ps] or [9]), f.get("firma", ""))
    P.append("<table class='small'><tr><th>Firma</th><th>Firma türü</th><th>Faaliyet</th><th>Paket(ler)</th><th>Ülke</th><th>Durum</th></tr>")
    for f in sorted(firms, key=key):
        P.append(f"<tr><td><b>{esc(f.get('firma'))}</b></td><td>{esc(f.get('firma_turu'))}</td><td>{esc(f.get('faaliyet'))}</td><td>{esc(f.get('paketler'))}</td><td>{esc(f.get('ulke'))}</td><td>{esc(f.get('durum'))}</td></tr>")
    P.append("</table></div>")

doc = f"<!doctype html><html lang='tr'><head><meta charset='utf-8'><title>URF 1. Etap MEP Bütçe Sunumu</title><style>{css}</style></head><body>{''.join(P)}</body></html>"
open(OUT_HTML, "w", encoding="utf-8").write(doc)
js = f"""const {{ chromium }} = require('playwright');
(async () => {{ const b = await chromium.launch(); const p = await b.newPage();
await p.goto('file://{os.path.abspath(OUT_HTML)}'); await p.pdf({{ path: '{os.path.abspath(OUT_PDF)}', format: 'A4', printBackground: true,
displayHeaderFooter: true, headerTemplate: '<span></span>', footerTemplate: '<div style="font-size:8px;width:100%;text-align:center;color:#9CA3AF">Şanlıurfa HRS 1. Etap – MEP Bütçe Sunumu · MEP Center · Sayfa <span class=pageNumber></span> / <span class=totalPages></span></div>',
margin: {{ top: '14mm', bottom: '16mm', left: '13mm', right: '13mm' }} }}); await b.close(); }})();"""
open("/tmp/_pdf.js", "w").write(js)
subprocess.run(["node", "/tmp/_pdf.js"], check=True, cwd=os.path.dirname(os.path.abspath(OUT_HTML)))
print("OK", OUT_PDF)
